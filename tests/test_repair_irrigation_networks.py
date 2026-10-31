"""Native-style regression, exact case transcription and display-only AOA invariants."""
from pathlib import Path
from zipfile import ZipFile
import io,json,math,random,unittest,xml.etree.ElementTree as E
import numpy as np
import pandas as pd
import openpyxl
import networkx as nx
from core.network import calculate_cpm,build_network
from core.reference_analysis import reference_cpm,network_svg_html
from core.simulation import simulate
from data_io.input_reader import read_v1_workbook,project_settings
from data_io.model_adapter import to_model_objects
from core.validation import validate_all
from localization import PT,set_language,network_view
from presentation.aoa import build_aoa,aoa_svg_html
ROOT=Path(__file__).resolve().parents[1]
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
Q=lambda n:'{'+NS['s']+'}'+n

def valid_dxf_order(raw):
    order={k:i for i,k in enumerate(['font','numFmt','fill','alignment','border','protection','extLst'])}
    tree=E.fromstring(raw)
    for d in tree.find(Q('dxfs')):
        tags=[c.tag.rsplit('}',1)[-1] for c in d]
        if any(k not in order for k in tags) or len(tags)!=len(set(tags)) or [order[k] for k in tags]!=sorted(order[k] for k in tags):return False
    return True

def signature(node):return (node.tag,sorted(node.attrib.items()),(node.text or '').strip(),[signature(c) for c in node])

class RepairTests(unittest.TestCase):
 def tearDown(self):set_language('en')
 def test_old_defect_is_rejected_and_original_styles_restored(self):
  with ZipFile(ROOT/'baseline/previous_distribution_release.zip') as z:
   with ZipFile(io.BytesIO(z.read('ProjectRisk_Distributions_EN_PT_BR/project_risk_template_v2_2.xlsx'))) as old:self.assertFalse(valid_dxf_order(old.read('xl/styles.xml')))
  with ZipFile(ROOT/'baseline/ProjectRisk_V2_2_original.zip') as z:
   with ZipFile(io.BytesIO(z.read('ProjectRisk_V2_2/project_risk_template_v2_2.xlsx'))) as original:styles=E.fromstring(original.read('xl/styles.xml'))
  for file in list(ROOT.glob('*.xlsx'))+list((ROOT/'examples').glob('*.xlsx')):
   with ZipFile(file) as z:
    raw=z.read('xl/styles.xml');self.assertTrue(valid_dxf_order(raw),file.name)
    now=E.fromstring(raw);self.assertEqual(signature(now.find(Q('dxfs'))),signature(styles.find(Q('dxfs'))))
    self.assertEqual(signature(now.find(Q('fills'))),signature(styles.find(Q('fills'))))
    for name in z.namelist():
     if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
      for rule in E.fromstring(z.read(name)).findall('.//s:cfRule',NS):
       if 'dxfId' in rule.attrib:self.assertLess(int(rule.get('dxfId')),len(now.find(Q('dxfs'))))
    calc=E.fromstring(z.read('xl/workbook.xml')).find(Q('calcPr'));self.assertEqual(calc.get('calcMode'),'auto')
    for name in z.namelist():
     if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
      cols=E.fromstring(z.read(name)).find(Q('cols'))
      if cols is not None:
       previous=0
       for col in cols:self.assertGreater(int(col.get('min')),previous);previous=int(col.get('max'))
 def test_parameter_highlights_keep_original_colors(self):
  for file in list(ROOT.glob('*.xlsx'))+list((ROOT/'examples').glob('*.xlsx')):
   w=openpyxl.load_workbook(file);pt='pt_BR' in file.name
   for worksheet in w:
    for table in worksheet.tables.values():
     a=openpyxl.utils.range_boundaries(table.ref)
     self.assertEqual([c.name for c in table.tableColumns],[str(worksheet.cell(a[1],j).value) for j in range(a[0],a[2]+1)])
   for sheet,target in [(('Atividades' if pt else 'Activities'),'FFFFF2CC'),(('Mapa_Risco_Atividade' if pt else 'Risk_Activity_Map'),'FFC6EFCE')]:
    for region in w[sheet].conditional_formatting:
     for rule in w[sheet].conditional_formatting[region]:
      if rule.type=='expression' and rule.formula and ('$D3=' in rule.formula[0] or '$N3=' in rule.formula[0]):
       self.assertEqual(rule.dxf.fill.bgColor.rgb,target,(file.name,str(region)))
 def test_source_values_and_all_user_decisions(self):
  expected=json.loads((ROOT/'verification/IRRIGATION_EXPECTED.json').read_text())
  for lang in ['en','pt_BR']:
   data=read_v1_workbook(ROOT/f'examples/irrigation_case_{lang}.xlsx');self.assertEqual(validate_all(data),[])
   a,r,m=to_model_objects(data);self.assertEqual((len(a),len(r),len(m)),(15,8,16))
   self.assertEqual([x['id'] for x in a],[x['id'] for x in expected['activities']])
   self.assertIn('2.1',[x['id'] for x in a]);self.assertIn('2.10',[x['id'] for x in a])
   cols=['Duration_Fixed','Duration_Min','Duration_MostLikely','Duration_Max','Duration_Mean','Duration_SD']
   for row,exp,act in zip(data['Activities'].to_dict('records'),expected['activities'],a):
    self.assertEqual(act['name'],exp['name_en' if lang=='en' else 'name_pt'])
    self.assertEqual(act['predecessors'],exp['predecessors']);self.assertEqual(act['duration_model'],exp['model'])
    for col,value in zip(cols,exp['parameters']):
     if value is None:self.assertTrue(pd.isna(row[col]))
     else:self.assertEqual(row[col],value)
   cols=[c.replace('Duration_','Impact_') for c in cols]
   for risk,exp in zip(r,expected['risks']):
    self.assertEqual(risk['id'],exp['id']);self.assertEqual(risk['probability'],exp['probability'])
    rows=data['Risk_Activity_Map'].loc[data['Risk_Activity_Map']['Risk_ID']==risk['id']]
    self.assertEqual(rows['Activity_ID'].tolist(),exp['affected'])
    self.assertEqual(rows['Distribution'].tolist(),[exp['model']]*2)
    for row in rows.to_dict('records'):
     for col,value in zip(cols,exp['parameters']):
      if value is None:self.assertTrue(pd.isna(row[col]))
      else:self.assertEqual(row[col],value)
   self.assertEqual([x['activity_id'] for x in m if x['risk_id']=='R8'],['2.4','2.12'])
   statement=data['Risk_Register'].loc[data['Risk_Register']['Risk_ID']=='R8','Risk_Statement'].iloc[0]
   self.assertIn('2.4',statement);self.assertNotIn('(2.15)',statement)
   self.assertEqual(next(x for x in m if x['risk_id']=='R3')['distribution'],'Normal')
   self.assertTrue(data['Risk_Relationships'].empty);self.assertTrue(data['Risk_Treatments'].empty)
   self.assertTrue(all(x['cost_params']=={'fixed':0.0} and x['cost_rate']==0 for x in a))
   settings=project_settings(data['Project']);self.assertEqual(settings['Time_Unit'],'weeks');self.assertTrue(settings['Input_Notes'])
 def test_irrigation_bilingual_results_and_cpm(self):
  objects=[to_model_objects(read_v1_workbook(ROOT/f'examples/irrigation_case_{lang}.xlsx')) for lang in ['en','pt_BR']]
  results=[simulate(*o,iterations=1000,seed=42) for o in objects]
  for key in results[0]:pd.testing.assert_frame_equal(results[0][key],results[1][key],check_exact=True)
  self.assertTrue((results[0]['project']['Cost']==0).all())
  a=objects[0][0];self.assertEqual(reference_cpm(a)[1]['project_duration'],26.5);self.assertEqual(build_aoa(a)['project_duration'],26.5)
 def test_aoa_preserves_precedence_and_all_cpm_fields(self):
  for seed in range(30):
   rng=random.Random(seed);acts=[];dur={}
   for i in range(15):
    aid=f'A{i}';preds=[f'A{j}' for j in range(i) if rng.random()<.2]
    acts.append({'id':aid,'predecessors':preds});dur[aid]=0 if i%7==0 else rng.uniform(.1,10)
   ordinary=calculate_cpm(acts,dur);aoa=build_aoa(acts,dur);g=aoa['graph'];dag=build_network(acts)
   self.assertTrue(nx.is_directed_acyclic_graph(g));self.assertEqual(sum(not e['dummy'] for _,_,e in g.edges(data=True)),len(acts))
   self.assertAlmostEqual(ordinary['project_duration'],aoa['project_duration'],places=10)
   for aid in dag:
    for field in ['ES','EF','LS','LF','Slack']:self.assertAlmostEqual(ordinary[field][aid],aoa['activities'][aid][field],places=10)
    self.assertEqual(aid in ordinary['critical_activities'],aoa['activities'][aid]['Critical'])
    for other in dag:
     if aid!=other:self.assertEqual(nx.has_path(dag,aid,other),nx.has_path(g,aoa['activities'][aid]['finish_event'],aoa['activities'][other]['start_event']))
 def test_numeric_identifiers_and_svg_are_safe(self):
  acts=to_model_objects(read_v1_workbook(ROOT/'examples/irrigation_case_en.xlsx'))[0]
  table,_=reference_cpm(acts);set_language('pt_BR')
  self.assertIn('>2.10</text>',network_view(network_svg_html(acts,table)))
  self.assertIn('>2.10</text>',network_view(aoa_svg_html(acts,lang='pt_BR')))
  self.assertNotIn('>2,10</text>',network_view(network_svg_html(acts,table)))
  svg=aoa_svg_html(acts);E.fromstring(svg)
  self.assertEqual(svg.count('data-activity='),15);self.assertIn('data-dummy="true"',svg)
 def test_display_conversion_does_not_mutate_inputs(self):
  import copy
  a=to_model_objects(read_v1_workbook(ROOT/'project_risk_template_v2_2.xlsx'))[0];before=copy.deepcopy(a)
  for language in ['en','pt_BR']:aoa_svg_html(a,lang=language)
  self.assertEqual(a,before)

if __name__=='__main__':unittest.main(verbosity=2)
