"""Boundary, locale, export and state checks; execute from the project directory."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import copy
import json
import sys
import tempfile
import unittest
import pandas as pd
import openpyxl
from localization import set_language,t,number,localized_frame,localized_csv,table_view
from localization.schema import normalize_table,canonical_header,resolve_sheets,CONTROLLED,numeric
from localization.exports import results_zip
from data_io.input_reader import read_v1_workbook, _read_table, HEADER_KEYS
from data_io.model_adapter import to_model_objects
from data_io.export_results import results_zip as original_zip
from core.validation import validate_all
from core.simulation import simulate

ROOT=Path(__file__).resolve().parents[1]
class BilingualTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.en=read_v1_workbook(ROOT/'project_risk_template_v2_2.xlsx')
        cls.pt=read_v1_workbook(ROOT/'project_risk_template_v2_2_pt_BR.xlsx')
    def tearDown(self):set_language('en')
    def test_model_equivalence(self):
        for a,b in zip(to_model_objects(self.en),to_model_objects(self.pt)):
            self.assertEqual([{k:v for k,v in x.items() if k!='name'} for x in a],[{k:v for k,v in x.items() if k!='name'} for x in b])
        self.assertEqual(validate_all(self.en),validate_all(self.pt))
    def test_all_controlled_values(self):
        from localization import PT
        for col,values in CONTROLLED.items():
            for value in values:
                frame=pd.DataFrame({col:[PT.get(value,value)]})
                self.assertEqual(normalize_table(frame,'Risk_Register',True).iloc[0,0],value)
    def test_boolean_inactive_validation(self):
        frame=pd.DataFrame({'Active':['não','falso','sim','verdadeiro',False,True]})
        self.assertEqual(normalize_table(frame,'Risk_Register',True)['Active'].tolist(),[False,False,True,True,False,True])
        data=copy.deepcopy(self.pt);data['Risk_Register'].loc[0,'Active']=False;data['Risk_Register'].loc[0,'Probability']=None
        self.assertFalse(any('Probability must' in x.message and 'R01:' in x.message for x in validate_all(data)))
    def test_headers_and_collisions(self):
        self.assertEqual(canonical_header(' ID_da_Atividade '),'Activity_ID')
        self.assertEqual(canonical_header('ID_do_Risco'),'Risk_ID')
        with self.assertRaises(ValueError):resolve_sheets(['Activities','Atividades'])
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'duplicate.xlsx'
            w=openpyxl.Workbook();s=w.active;s.title='Activities';s.append(['Activity_ID','Duration_Model','ID_da_Atividade']);w.save(path)
            with self.assertRaisesRegex(ValueError,'duplicate column'):_read_table(path,'Activities')
    def test_locale_numbers(self):
        from localization import display_value
        set_language('pt_BR');self.assertNotEqual(display_value(0.0000001),'0,00');self.assertEqual(number(1234.56),'1.234,56');self.assertEqual(number(.255,'.1%'),'25,5%')
        self.assertEqual(numeric('1.234,56',True),1234.56);self.assertEqual(numeric('1,234.56',False),1234.56)
        with self.assertRaises(ValueError):numeric('1,234.56',True)
        with self.assertRaises(ValueError):numeric('1.234,56',False)
        self.assertEqual(numeric(1.23456,True),1.23456)
    def test_validation_translation(self):
        set_language('pt_BR')
        fixtures=[('A01 references unknown predecessor A99.','A01 faz referência à predecessora inexistente A99.'),('R01: Probability must be between 0 and 1.','R01: Probabilidade deve estar entre 0 e 1.')]
        for source,target in fixtures:self.assertEqual(t(source),target)
        for source in ["Worksheet 'Activities': duplicate column names in header row 2.",'Missing required worksheet(s): Project, Activities','R01: acceptance documentation incomplete: Residual_Probability, Residual_Rationale']:
            self.assertNotEqual(t(source),source)
    def test_display_does_not_mutate(self):
        df=pd.DataFrame({'Activity_ID':['Threat'],'Duration':[1234.56],'Risk_Name':['Threat']})
        saved=df.copy(deep=True);set_language('pt_BR')
        display=localized_frame(df);table_view(df)
        pd.testing.assert_frame_equal(df,saved)
        self.assertEqual(display['Nome_do_Risco'][0],'Threat')
        self.assertEqual(display['Duração'][0],1234.56)
        self.assertEqual(display['ID_da_Atividade'][0],'Threat')
        self.assertIn('Threat',t('Treated probability — Threat'))
    def test_exports_keep_original_bytes(self):
        frame=pd.DataFrame({'Risk_ID':['R01'],'Duration':[1234.56],'Critical':[True]})
        set_language('pt_BR')
        old=original_zip(*([frame]*5));new=results_zip(*([frame]*5))
        with ZipFile(BytesIO(old)) as a,ZipFile(BytesIO(new)) as b:
            for name in a.namelist():self.assertEqual(a.read(name),b.read(name))
            raw=b.read('pt_BR/project_iterations.csv').decode('utf-8-sig')
            self.assertIn('1234,56',raw);self.assertIn(';',raw);self.assertIn('verdadeiro',raw)
        restored=pd.read_csv(BytesIO(localized_csv(frame)),sep=';',decimal=',')
        self.assertEqual(restored.iloc[0,1],1234.56)
    def test_template_native_preservation(self):
        a=openpyxl.load_workbook(ROOT/'project_risk_template_v2_2.xlsx',data_only=False)
        b=openpyxl.load_workbook(ROOT/'project_risk_template_v2_2_pt_BR.xlsx',data_only=False)
        self.assertEqual(len(a.sheetnames),len(b.sheetnames))
        from localization import PT
        for sa,sb in zip(a,b):
            self.assertEqual((sa.max_row,sa.max_column),(sb.max_row,sb.max_column))
            self.assertEqual(str(sa.freeze_panes),str(sb.freeze_panes))
            self.assertEqual(set(map(str,sa.merged_cells.ranges)),set(map(str,sb.merged_cells.ranges)))
            self.assertEqual(len(sa.data_validations.dataValidation),len(sb.data_validations.dataValidation))
            self.assertEqual(len(sa.conditional_formatting),len(sb.conditional_formatting))
            for row in sa:
                for cell in row:
                    target=sb[cell.coordinate]
                    if isinstance(cell.value,(int,float,bool)):
                        self.assertEqual((cell.value,cell.data_type),(target.value,target.data_type))
                    if cell.data_type=='f':self.assertEqual(target.data_type,'f')
        with ZipFile(ROOT/'project_risk_template_v2_2.xlsx') as za,ZipFile(ROOT/'project_risk_template_v2_2_pt_BR.xlsx') as zb:
            self.assertEqual(za.namelist(),zb.namelist())
    def test_seeded_language_equivalence(self):
        a=simulate(*to_model_objects(self.en),iterations=100,seed=42)
        b=simulate(*to_model_objects(self.pt),iterations=100,seed=42)
        for key in a:pd.testing.assert_frame_equal(a[key],b[key],check_exact=True)

if __name__=='__main__':unittest.main(verbosity=2)
